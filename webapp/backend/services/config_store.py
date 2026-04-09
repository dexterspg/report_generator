"""
Config Store — Slice 2 (Account Mapping only)

Handles:
  - SQLite database initialisation (_init_db)
  - Path resolution (_get_app_data_dir)
  - Account mapping CRUD: get, save (replace/merge), reset, export
  - File parsing with flexible column alias matching

Exchange rates, history, and rate-lookup queries are added in Slices 3 and 5.
"""

import io
import os
import sqlite3
import sys
import uuid
from pathlib import Path

import pandas as pd
from fastapi import UploadFile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ---------------------------------------------------------------------------
# Column alias definitions — keys are canonical names, values are accepted
# lowercase variants the user might put in their spreadsheet.
# ---------------------------------------------------------------------------

ACCOUNT_MAPPING_ALIASES: dict[str, list[str]] = {
    "account_number": [
        "account number",
        "account_number",
        "acct number",
        "acct_number",
        "acct no",
        "acct no.",
        "acct num",
        "account no",
        "account no.",
        "account#",
        "gl account",
        "gl_account",
        "glaccount",
    ],
    "account_type": [
        "account type",
        "account_type",
        "acct type",
        "acct_type",
        "type",
    ],
    "monetary": [
        "monetary?",
        "monetary",
        "is_monetary",
        "is monetary",
        "monetary classification",
        "monetary class",
    ],
    "rate": [
        "rate",
        "rate method",
        "rate_method",
        "rate type",
        "rate_type",
        "remeasurement rate",
    ],
}

# Canonical display-name order for the downloaded Excel
ACCOUNT_MAPPING_HEADERS = ["Account Number", "Account Type", "Monetary?", "Rate"]


# ---------------------------------------------------------------------------
# Path resolution
# ---------------------------------------------------------------------------

def _get_app_data_dir() -> Path:
    """
    Resolve the application data directory.

    Desktop mode  (.exe via PyInstaller) → %LOCALAPPDATA%/CTR-FX-Remeasurement/
    Development mode (python app.py)     → webapp/backend/  (same folder as app.py)
    """
    if getattr(sys, "frozen", False):
        # PyInstaller .exe — store data in %LOCALAPPDATA%
        local_app_data = os.environ.get("LOCALAPPDATA", str(Path.home()))
        data_dir = Path(local_app_data) / "CTR-FX-Remeasurement"
    else:
        # Development — store alongside app.py
        data_dir = Path(__file__).parent.parent  # webapp/backend/

    data_dir.mkdir(parents=True, exist_ok=True)
    return data_dir


def _get_db_path() -> Path:
    return _get_app_data_dir() / "ctr_fx.db"


# ---------------------------------------------------------------------------
# Database initialisation
# ---------------------------------------------------------------------------

def _get_connection() -> sqlite3.Connection:
    """Open (and if needed initialise) the SQLite database."""
    db_path = _get_db_path()
    conn = sqlite3.connect(str(db_path))
    conn.row_factory = sqlite3.Row
    _init_db(conn)
    return conn


def _init_db(conn: sqlite3.Connection) -> None:
    """Create tables if they don't exist yet. Safe to call on every connection."""
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS account_mapping (
            id             TEXT PRIMARY KEY,
            account_number TEXT NOT NULL UNIQUE,
            account_type   TEXT NOT NULL,
            monetary       TEXT NOT NULL,
            rate           TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS exchange_rates (
            id            TEXT PRIMARY KEY,
            rate_type     TEXT,
            from_currency TEXT NOT NULL,
            to_currency   TEXT NOT NULL,
            valid_from    TEXT NOT NULL,
            exchange_rate TEXT NOT NULL,
            UNIQUE(from_currency, to_currency, valid_from)
        );

        CREATE INDEX IF NOT EXISTS idx_exchange_rates_lookup
            ON exchange_rates (from_currency, to_currency, valid_from);

        CREATE TABLE IF NOT EXISTS config_history (
            id                        TEXT PRIMARY KEY,
            timestamp                 TEXT NOT NULL,
            action                    TEXT NOT NULL,
            config_type               TEXT NOT NULL,
            source_filename           TEXT,
            details                   TEXT,
            snapshot_account_mapping  TEXT,
            snapshot_exchange_rates   TEXT
        );
        """
    )
    conn.commit()


# ---------------------------------------------------------------------------
# Account Mapping — read
# ---------------------------------------------------------------------------

def get_account_mapping() -> list[dict]:
    """Return all saved account mapping rows as a list of dicts."""
    with _get_connection() as conn:
        rows = conn.execute(
            "SELECT account_number, account_type, monetary, rate "
            "FROM account_mapping ORDER BY account_number"
        ).fetchall()
    return [dict(r) for r in rows]


def get_account_mapping_count() -> int:
    """Return the number of saved account mapping entries."""
    with _get_connection() as conn:
        row = conn.execute("SELECT COUNT(*) as cnt FROM account_mapping").fetchone()
    return row["cnt"]


# ---------------------------------------------------------------------------
# Account Mapping — parse uploaded file
# ---------------------------------------------------------------------------

def _normalise_col(name: str) -> str:
    """Lowercase + strip whitespace for flexible column matching."""
    return str(name).lower().strip()


def _resolve_column(df_columns: list[str], canonical: str) -> str | None:
    """
    Return the actual DataFrame column name that maps to `canonical`.
    Returns None if no alias matches.
    """
    aliases = ACCOUNT_MAPPING_ALIASES[canonical]
    for col in df_columns:
        if _normalise_col(col) in aliases:
            return col
    return None


def parse_account_mapping_file(upload: UploadFile) -> list[dict]:
    """
    Parse an uploaded .xlsx / .xls / .csv file.

    Returns a list of dicts with canonical keys:
        account_number, account_type, monetary, rate

    Raises ValueError with a clear message if required columns are missing.
    """
    filename = upload.filename or ""
    content = upload.file.read()

    if filename.lower().endswith(".csv"):
        df = pd.read_csv(io.BytesIO(content), dtype=str)
    else:
        df = pd.read_excel(io.BytesIO(content), dtype=str)

    # Strip whitespace from all column names
    df.columns = [str(c).strip() for c in df.columns]

    # Resolve columns using alias matching
    col_map: dict[str, str] = {}
    missing: list[str] = []

    for canonical in ACCOUNT_MAPPING_ALIASES:
        resolved = _resolve_column(list(df.columns), canonical)
        if resolved is None:
            missing.append(canonical.replace("_", " ").title())
        else:
            col_map[canonical] = resolved

    if missing:
        raise ValueError(
            f"Missing required column(s): {', '.join(missing)}. "
            f"Columns found in file: {', '.join(df.columns.tolist())}."
        )

    # Build normalised records
    records: list[dict] = []
    for _, row in df.iterrows():
        acct_num = str(row[col_map["account_number"]]).strip()
        if not acct_num or acct_num.lower() == "nan":
            continue  # skip blank rows

        records.append(
            {
                "account_number": acct_num,
                "account_type": str(row[col_map["account_type"]]).strip(),
                "monetary": str(row[col_map["monetary"]]).strip(),
                "rate": str(row[col_map["rate"]]).strip(),
            }
        )

    if not records:
        raise ValueError(
            "No data rows found in the uploaded file. "
            "Verify the file has data rows below the header."
        )

    return records


# ---------------------------------------------------------------------------
# Account Mapping — save (replace mode)
# ---------------------------------------------------------------------------

def save_account_mapping_replace(records: list[dict]) -> int:
    """
    Replace mode: clear all existing entries, then bulk-insert from records.
    Returns the number of rows inserted.
    """
    with _get_connection() as conn:
        conn.execute("DELETE FROM account_mapping")
        _bulk_insert_account_mapping(conn, records)
        conn.commit()
    return len(records)


# ---------------------------------------------------------------------------
# Account Mapping — merge mode
# ---------------------------------------------------------------------------

def save_account_mapping_merge(records: list[dict]) -> int:
    """
    Merge mode: INSERT OR REPLACE using the UNIQUE constraint on account_number.
    New entries are added; existing entries (same account_number) are updated.
    Returns the number of rows processed from the file.
    """
    with _get_connection() as conn:
        _bulk_insert_account_mapping(conn, records)
        conn.commit()
    return len(records)


def _bulk_insert_account_mapping(conn: sqlite3.Connection, records: list[dict]) -> None:
    """INSERT OR REPLACE all records. Relies on UNIQUE(account_number)."""
    conn.executemany(
        """
        INSERT OR REPLACE INTO account_mapping (id, account_number, account_type, monetary, rate)
        VALUES (?, ?, ?, ?, ?)
        """,
        [
            (
                str(uuid.uuid4()),
                r["account_number"],
                r["account_type"],
                r["monetary"],
                r["rate"],
            )
            for r in records
        ],
    )


# ---------------------------------------------------------------------------
# Account Mapping — reset
# ---------------------------------------------------------------------------

def reset_account_mapping() -> None:
    """Delete all saved account mapping entries."""
    with _get_connection() as conn:
        conn.execute("DELETE FROM account_mapping")
        conn.commit()


# ---------------------------------------------------------------------------
# Account Mapping — export as styled Excel
# ---------------------------------------------------------------------------

def export_account_mapping() -> bytes:
    """
    Export the current account mapping as a styled .xlsx file.
    The produced file is a valid upload file (closes the round-trip / sharing loop).

    Returns the raw bytes of the workbook.
    """
    rows = get_account_mapping()

    wb = Workbook()
    ws = wb.active
    ws.title = "Account Mapping"

    # ---- Styles ----
    header_font = Font(name="Calibri", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="009CDE")  # Nakisa brand blue
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
    cell_alignment = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        bottom=Side(style="thin", color="E4E7EC"),
    )

    # ---- Header row ----
    for col_idx, header in enumerate(ACCOUNT_MAPPING_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment

    # ---- Data rows ----
    for row_idx, record in enumerate(rows, start=2):
        values = [
            record["account_number"],
            record["account_type"],
            record["monetary"],
            record["rate"],
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = cell_alignment
            cell.border = thin_border

        # Zebra striping on even data rows
        if row_idx % 2 == 0:
            for col_idx in range(1, 5):
                ws.cell(row=row_idx, column=col_idx).fill = PatternFill(
                    "solid", fgColor="F7F8FA"
                )

    # ---- Column widths ----
    col_widths = [20, 22, 18, 16]
    for col_idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[
            ws.cell(row=1, column=col_idx).column_letter
        ].width = width

    # ---- Freeze header row ----
    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
